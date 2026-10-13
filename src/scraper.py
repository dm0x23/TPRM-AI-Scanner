#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import difflib
import hashlib
import html as htmllib
import json
import os
import re
import sys
import time
from datetime import datetime, timezone
from urllib import robotparser
from urllib.parse import urljoin, urlparse, urldefrag, quote_plus

import requests
from bs4 import BeautifulSoup, NavigableString
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

# --------------------------------------------------------------------------- #
# 1. LEVEL CONFIG  (agree the number of levels with the Task-1 owner; edit here
#    or pass --levels-config levels.json with the same structure)
#    pages = max pages fetched per source type at that level
# --------------------------------------------------------------------------- #
DEFAULT_LEVELS = {
    1: {"label": "Light (Low)",
        "pages": {"homepage": 1, "product": 3, "trust_privacy": 3},
        "probe": ["trust_privacy"], "github": False, "edgar": False, "news": False},
    2: {"label": "Standard (Medium)",
        "pages": {"homepage": 1, "product": 5, "trust_privacy": 4, "careers": 3, "news_press": 4},
        "probe": ["trust_privacy", "careers", "news_press"], "github": False, "edgar": False, "news": False},
    3: {"label": "Deep (High)",
        "pages": {"homepage": 1, "product": 6, "trust_privacy": 5, "careers": 5, "news_press": 6, "docs_dev": 4},
        "probe": ["trust_privacy", "careers", "news_press", "docs_dev"], "github": True, "edgar": True, "news": False},
    4: {"label": "Full (Critical)",
        "pages": {"homepage": 1, "product": 10, "trust_privacy": 8, "careers": 8, "news_press": 10, "docs_dev": 8},
        "probe": ["trust_privacy", "careers", "news_press", "docs_dev"], "github": True, "edgar": True, "news": True},
}
TIER_TO_LEVEL = {"low": 1, "medium": 2, "high": 3, "critical": 4}

MAX_EVIDENCE_PER_PAGE = 12
CTRL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")  # break Excel/CSV exports
SKIP_EXT = re.compile(r"\.(pdf|jpe?g|png|gif|svg|webp|ico|zip|gz|mp[34]|css|js|woff2?|docx?|xlsx?|pptx?)$", re.I)
MAX_SENTENCE_CHARS = 700

# --------------------------------------------------------------------------- #
# 2. KEYWORDS
# --------------------------------------------------------------------------- #
AI_TERMS = [  # (label, compiled regex)
    ("AI", re.compile(r"\bAI\b")),
    ("artificial intelligence", re.compile(r"artificial[\s-]intelligence", re.I)),
    ("machine learning", re.compile(r"machine[\s-]learning|\bML\b", re.I)),
    ("deep learning", re.compile(r"deep[\s-]learning|neural[\s-]network", re.I)),
    ("LLM", re.compile(r"\bLLMs?\b|large[\s-]language[\s-]model", re.I)),
    ("generative AI", re.compile(r"generative|\bgen[\s-]?AI\b", re.I)),
    ("foundation model", re.compile(r"foundation[\s-]model", re.I)),
    ("NLP", re.compile(r"natural[\s-]language[\s-]processing|\bNLP\b", re.I)),
    ("computer vision", re.compile(r"computer[\s-]vision", re.I)),
    ("agentic", re.compile(r"agentic|\bAI agents?\b|autonomous agents?", re.I)),
    ("copilot/assistant", re.compile(r"copilot|\bchat-?bot\b|virtual assistant", re.I)),
    ("RAG/embeddings", re.compile(r"retrieval[\s-]augmented|\bRAG\b|vector (database|store)|embeddings?", re.I)),
    ("predictive model", re.compile(r"predictive (model|analytics)|anomaly detection|model training", re.I)),
]
PROVIDERS = [  # (canonical name, regex)  -> third-party model / AI platform names
    ("OpenAI", re.compile(r"\bOpenAI\b|\bChatGPT\b|\bGPT-?\d", re.I)),
    ("Azure OpenAI", re.compile(r"Azure OpenAI|Azure AI", re.I)),
    ("Anthropic/Claude", re.compile(r"\bAnthropic\b|\bClaude\b")),
    ("Google (Gemini/Vertex)", re.compile(r"\bGemini\b|Vertex AI|Google Cloud AI|\bPaLM\b")),
    ("AWS (Bedrock/SageMaker/Comprehend/Textract)", re.compile(r"Bedrock|SageMaker|Amazon Comprehend|Amazon Textract|Amazon Q\b")),
    ("Microsoft Copilot", re.compile(r"Microsoft (365 )?Copilot|GitHub Copilot", re.I)),
    ("Cohere", re.compile(r"\bCohere\b")),
    ("Mistral", re.compile(r"\bMistral\b")),
    ("Meta Llama", re.compile(r"\bLlama[\s-]?\d?", re.I)),
    ("Hugging Face", re.compile(r"Hugging ?Face", re.I)),
    ("Databricks", re.compile(r"Databricks|MosaicML")),
    ("IBM watsonx", re.compile(r"watsonx", re.I)),
    ("NVIDIA", re.compile(r"\bNVIDIA\b")),
    ("Salesforce Einstein", re.compile(r"Einstein (GPT|AI|Copilot)", re.I)),
    ("ServiceNow Now Assist", re.compile(r"Now Assist")),
    ("Palantir AIP", re.compile(r"Palantir AIP")),
    ("DataRobot/H2O", re.compile(r"DataRobot|H2O\.ai", re.I)),
]

# URL-pattern routing: source type -> keywords expected in host/path
TYPE_PATTERNS = {
    "trust_privacy": ["trust", "security", "privacy", "subprocessor", "sub-processor", "responsible-ai",
                      "ai-policy", "ai-governance", "governance", "compliance", "legal", "dpa", "data-protection"],
    "product": ["product", "solution", "platform", "service", "capabilit", "offering", "what-we-do",
                "automation", "intelligence", "analytics", "fraud", "/ai", "-ai", "innovation"],
    "careers": ["career", "jobs", "join-us", "work-with-us", "opportunit", "talent"],
    "news_press": ["news", "press", "blog", "insight", "media", "article", "newsroom", "announce", "thought-leadership"],
    "docs_dev": ["docs", "developer", "api", "integration", "documentation", "knowledge", "support", "release-notes"],
}
# Well-known paths probed on the vendor domain (negative findings get logged)
PROBE_PATHS = {
    "trust_privacy": ["/trust", "/security", "/trust-center", "/privacy", "/privacy-policy", "/privacy-notice",
                      "/subprocessors", "/sub-processors", "/legal/subprocessors", "/responsible-ai", "/ai-policy"],
    "careers": ["/careers", "/jobs", "/about/careers", "/company/careers"],
    "news_press": ["/newsroom", "/news", "/press", "/blog", "/insights", "/about/news"],
    "docs_dev": ["/developers", "/developer", "/docs", "/api"],
}
ATS_PATTERNS = {
    "greenhouse": re.compile(r"(?:boards|job-boards)\.greenhouse\.io/(?:embed/job_board\?for=)?([\w-]+)", re.I),
    "lever": re.compile(r"jobs\.lever\.co/([\w-]+)", re.I),
    "ashby": re.compile(r"jobs\.ashbyhq\.com/([\w.-]+)", re.I),
    "workday": re.compile(r"([\w-]+)\.wd\d+\.myworkdayjobs\.com", re.I),
    "smartrecruiters": re.compile(r"careers\.smartrecruiters\.com/([\w-]+)", re.I),
}
# vendor-specific lookups the heuristics cannot guess (extend / use --overrides)
NAME_OVERRIDES = {"bny": {"sec_name": "BANK OF NEW YORK MELLON CORP"}}


def now_iso() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def reg_domain(host: str) -> str:
    host = host.lower().split(":")[0]
    if re.fullmatch(r"[\d.]+", host):
        return host
    p = host.split(".")
    if len(p) >= 3 and p[-2] in {"co", "com", "org", "net", "gov", "ac"} and len(p[-1]) == 2:
        return ".".join(p[-3:])
    return ".".join(p[-2:])


def norm_url(site: str) -> str:
    site = site.strip()
    if not re.match(r"^https?://", site, re.I):
        site = "https://" + site
    return site.rstrip("/")


def origin_of(url: str) -> str:
    u = urlparse(url)
    return f"{u.scheme}://{u.netloc}"


def tier_to_level(tier):
    """'Critical' / 'High' / 'Tier 3' / 3  ->  1-4, else None"""
    t = str(tier or "").strip().lower()
    if t.isdigit():
        return int(t) if int(t) in DEFAULT_LEVELS else None
    for word in ("critical", "high", "medium", "low"):
        if word in t:
            return TIER_TO_LEVEL[word]
    m = re.search(r"\b([1-4])\b", t)
    return int(m.group(1)) if m else None


def sec_name_override(name: str):
    n = clean_name(name)
    for k, v in NAME_OVERRIDES.items():
        if n == k or n.startswith(k + " "):
            return v["sec_name"]
    return None


def json_or_none(text: str):
    try:
        return json.loads(text)
    except ValueError:
        return None


def clean_name(n: str) -> str:
    n = n or ""
    n = re.sub(r"\(.*?\)", " ", n.lower())
    n = re.sub(r"\b(inc|corp|corporation|co|company|llc|l\.l\.c|ltd|plc|the|group|holdings|technologies)\b\.?", " ", n)
    return " ".join(re.sub(r"[^a-z0-9 ]", " ", n).split())


# --------------------------------------------------------------------------- #
# 3. POLITE FETCHER
# --------------------------------------------------------------------------- #
class Fetched:
    def __init__(self, url, status, text="", ctype="", note=""):
        self.url, self.status, self.text, self.ctype, self.note = url, status, text, ctype, note

    @property
    def ok(self):
        return self.status == 200 and bool(self.text)


class Fetcher:
    def __init__(self, contact: str | None = None, delay: float = 1.0, timeout: int = 15,
                 respect_robots: bool = True):
        ua = "VendorAIOSINT/1.0 (academic third-party-risk case study" + (f"; contact: {contact})" if contact else ")")
        self.s = requests.Session()
        self.s.headers.update({"User-Agent": ua, "Accept-Language": "en"})
        retry_kw = dict(total=2, backoff_factor=1, status_forcelist=(429, 500, 502, 503, 504),
                        respect_retry_after_header=True)
        try:
            retry = Retry(allowed_methods=("GET",), **retry_kw)
        except TypeError:  # urllib3 < 1.26 calls it method_whitelist
            retry = Retry(method_whitelist=("GET",), **retry_kw)
        self.s.mount("http://", HTTPAdapter(max_retries=retry))
        self.s.mount("https://", HTTPAdapter(max_retries=retry))
        self.delay, self.timeout, self.respect_robots = delay, timeout, respect_robots
        self._last: dict[str, float] = {}
        self._robots: dict[str, robotparser.RobotFileParser | None] = {}
        self._cache: dict[str, Fetched] = {}
        self.log: list[dict] = []

    def _wait(self, host):
        gap = time.time() - self._last.get(host, 0)
        if gap < self.delay:
            time.sleep(self.delay - gap)
        self._last[host] = time.time()

    def robots_allows(self, url: str) -> bool:
        if not self.respect_robots:
            return True
        u = urlparse(url)
        base = f"{u.scheme}://{u.netloc}"
        if base not in self._robots:
            rp = robotparser.RobotFileParser()
            try:
                self._wait(u.netloc)
                r = self.s.get(base + "/robots.txt", timeout=self.timeout)
                rp.parse(r.text.splitlines() if r.status_code == 200 else [])
                self._robots[base] = rp
            except requests.RequestException:
                self._robots[base] = None
        rp = self._robots[base]
        return True if rp is None else rp.can_fetch(self.s.headers["User-Agent"], url)

    def robots_sitemaps(self, base: str) -> list[str]:
        rp = self._robots.get(base)
        return list(rp.site_maps() or []) if rp is not None and hasattr(rp, "site_maps") else []

    def get(self, url: str, *, robots: bool = True, max_bytes: int = 3_000_000, headers=None) -> Fetched:
        url, _ = urldefrag(url)
        if url in self._cache:
            return self._cache[url]
        if robots and not self.robots_allows(url):
            f = Fetched(url, 0, note="blocked_by_robots_txt")
        else:
            try:
                self._wait(urlparse(url).netloc)
                r = self.s.get(url, timeout=self.timeout, stream=True, headers=headers or {})
                try:
                    ctype = r.headers.get("Content-Type", "")
                    if r.status_code == 200 and not re.search(r"html|xml|json|text", ctype, re.I):
                        f = Fetched(r.url, 200, "", ctype, note="skipped_non_text_content")
                    else:
                        buf, size = [], 0
                        for chunk in r.iter_content(65536):
                            buf.append(chunk)
                            size += len(chunk)
                            if size > max_bytes:
                                break
                        enc = r.encoding
                        if not enc or (enc.lower() == "iso-8859-1" and "charset" not in ctype.lower()):
                            enc = "utf-8"  # requests guesses latin-1 for text/* with no charset -> garbled text
                        raw = b"".join(buf)
                        try:
                            body = raw.decode(enc, errors="replace")
                        except LookupError:
                            body = raw.decode("utf-8", errors="replace")
                        f = Fetched(r.url, r.status_code, body, ctype)
                finally:
                    r.close()  # stream=True: always release the connection
            except (requests.RequestException, ValueError) as e:
                f = Fetched(url, -1, note=f"request_error:{type(e).__name__}")
        self._cache[url] = f
        self.log.append({"url": url, "status": f.status, "note": f.note, "at": now_iso()})
        return f


# --------------------------------------------------------------------------- #
# 4. TEXT EXTRACTION + AI SIGNAL DETECTION
# --------------------------------------------------------------------------- #
def html_to_units(raw: str):
    """-> (title, published_date, list of sentence-level text units)"""
    soup = BeautifulSoup(raw, "html.parser")
    title = soup.title.get_text(" ", strip=True) if soup.title else ""
    pub = ""
    for sel in ({"property": "article:published_time"}, {"name": "date"}, {"itemprop": "datePublished"}):
        m = soup.find("meta", attrs=sel)
        if m and m.get("content"):
            pub = m["content"][:10]
            break
    if not pub:
        t = soup.find("time", attrs={"datetime": True})
        pub = t["datetime"][:10] if t else ""
    for t in soup(["script", "style", "noscript", "svg", "nav", "footer", "form", "iframe"]):
        t.decompose()
    units = []
    for line in soup.get_text("\n").split("\n"):
        line = re.sub(r"\s+", " ", line).strip()
        if len(line) < 4:
            continue
        for s in re.split(r"(?<=[.!?])\s+(?=[A-Z0-9\"“(])", line):
            s = s.strip()
            if len(s) >= 4:
                units.append(s)
    return title, pub, units


def find_signals(units: list[str], cap: int = MAX_EVIDENCE_PER_PAGE, min_len: int = 15):
    hits = []
    for i, u in enumerate(units):
        if len(u) < min_len:  # drops nav/link-text fragments like "Fraud AI"
            continue
        terms = [lbl for lbl, rx in AI_TERMS if rx.search(u)]
        provs = [name for name, rx in PROVIDERS if rx.search(u)]
        if not terms and not provs:
            continue
        ctx = " ".join(units[max(0, i - 1): i + 2])
        hits.append({"excerpt": u[:MAX_SENTENCE_CHARS], "context": ctx[:1200],
                     "matched_terms": terms, "named_providers": provs,
                     "_score": len(terms) + 3 * len(provs)})
    hits.sort(key=lambda h: -h["_score"])
    for h in hits:
        h.pop("_score")
    return hits[:cap]


# --------------------------------------------------------------------------- #
# 5. RESULT CONTAINER
# --------------------------------------------------------------------------- #
class Result:
    def __init__(self, vid, name, site, level, level_label):
        self.meta = {"vendor_id": vid, "vendor_name": name, "website": site, "level": level,
                     "level_label": level_label, "started_at": now_iso()}
        self.evidence, self.checked, self._seen = [], [], set()

    def add_evidence(self, *, source_type, origin, url, raw_for_hash, title="", published="", hit):
        key = (url, hit["excerpt"])
        if key in self._seen:
            return
        self._seen.add(key)
        rec = {
            "vendor_id": self.meta["vendor_id"], "vendor_name": self.meta["vendor_name"],
            "source_type": source_type, "origin": origin,  # first_party | third_party
            "source_url": url, "page_title": title, "published_date": published,
            "retrieved_at": now_iso(),
            "content_sha256": hashlib.sha256(raw_for_hash.encode("utf-8", "ignore")).hexdigest(),
            **hit}
        self.evidence.append({k: (CTRL_CHARS.sub("", v) if isinstance(v, str) else v) for k, v in rec.items()})

    def note(self, source_type, target, outcome, detail=""):
        self.checked.append({"source_type": source_type, "target": target, "outcome": outcome,
                             "detail": detail, "at": now_iso()})


# --------------------------------------------------------------------------- #
# 6. SOURCE COLLECTORS
# --------------------------------------------------------------------------- #
def classify_url(url: str):
    u = urlparse(url)
    s = (u.netloc.split(".")[0] + u.path).lower()
    best, best_score = None, 0
    for typ, kws in TYPE_PATTERNS.items():
        score = sum(1 for k in kws if k in s)
        if score > best_score:
            best, best_score = typ, score
    if not best:
        return None, 0
    depth = len([p for p in u.path.split("/") if p])
    return best, best_score * 10 - depth - len(u.path) / 80


def discover_urls(f: Fetcher, base: str, home: Fetched):
    site_dom = reg_domain(urlparse(base).netloc)
    origin = origin_of(base)  # sitemap / robots live at the site root even if base has a path
    urls = set()
    if home.ok:
        for a in BeautifulSoup(home.text, "html.parser").find_all("a", href=True):
            absu = urldefrag(urljoin(home.url, a["href"]))[0]
            if (urlparse(absu).scheme in ("http", "https") and reg_domain(urlparse(absu).netloc) == site_dom
                    and not SKIP_EXT.search(urlparse(absu).path)):
                urls.add(absu)
    sitemaps = f.robots_sitemaps(origin) or [origin + "/sitemap.xml"]
    queue, seen_sm = list(sitemaps), set()
    while queue and len(seen_sm) < 6 and len(urls) < 5000:
        sm = queue.pop(0)
        if sm in seen_sm:
            continue
        seen_sm.add(sm)
        r = f.get(sm)
        if not r.ok:
            continue
        for loc in re.findall(r"<loc>\s*(.*?)\s*</loc>", r.text, re.S):
            loc = htmllib.unescape(loc)
            if loc.endswith(".xml") or loc.endswith(".xml.gz"):
                queue.append(loc)
            elif reg_domain(urlparse(loc).netloc) == site_dom and not SKIP_EXT.search(urlparse(loc).path):
                urls.add(loc)
    return urls


def scan_page(f: Fetcher, res: Result, url: str, source_type: str, origin="first_party"):
    r = f.get(url)
    if not r.ok:
        return r, 0
    if "html" not in r.ctype.lower() and "<html" not in r.text[:500].lower():
        title, pub, units = "", "", [re.sub(r"\s+", " ", l).strip() for l in r.text.splitlines() if l.strip()]
    else:
        title, pub, units = html_to_units(r.text)
    n = 0
    for hit in find_signals(units):
        res.add_evidence(source_type=source_type, origin=origin, url=r.url, raw_for_hash=r.text,
                         title=title, published=pub, hit=hit)
        n += 1
    if sum(len(u) for u in units) < 200:
        res.note(source_type, r.url, "thin_content", "page may be JavaScript-rendered; consider Playwright fallback")
    return r, n


def collect_first_party(f, res, base, cfg):
    home, n = scan_page(f, res, base, "homepage")
    res.note("homepage", base, "hits" if n else ("no_hits" if home.ok else "unreachable"), home.note or str(home.status))
    discovered = discover_urls(f, base, home)
    res.note("discovery", base, "ok", f"{len(discovered)} same-site URLs found via homepage links + sitemap")

    # candidate pool per type: discovered URLs ranked + probe paths for well-known pages
    pool: dict[str, list[str]] = {t: [] for t in TYPE_PATTERNS}
    ranked = sorted(((u, *classify_url(u)) for u in discovered), key=lambda x: (-x[2], x[0]))  # tie-break = repeatable
    for u, typ, _ in ranked:
        if typ:
            pool[typ].append(u)
    for typ in cfg.get("probe", []):
        for p in PROBE_PATHS.get(typ, []):
            pool[typ].append(origin_of(base) + p)

    ats_found = {}
    for typ, limit in cfg["pages"].items():
        if typ == "homepage":
            continue
        fetched_ok, tried, seen = 0, 0, set()
        for u in pool.get(typ, []):
            if u in seen:
                continue
            seen.add(u)
            if fetched_ok >= limit or tried >= limit * 3:
                break
            tried += 1
            r, n = scan_page(f, res, u, typ)
            if r.ok:
                fetched_ok += 1
                res.note(typ, r.url, "hits" if n else "no_hits")
                if typ == "careers":
                    for ats, rx in ATS_PATTERNS.items():
                        m = rx.search(r.text)
                        if m:
                            ats_found[ats] = m.group(1)
            elif r.status in (404, 410):
                res.note(typ, u, "not_found", r.note or str(r.status))
            elif r.status == 0:
                res.note(typ, u, "blocked_by_robots", r.note)
            elif r.status == 200:  # reachable but not readable text (PDF, image, empty body)
                res.note(typ, u, "skipped", r.note or "empty_body")
            else:  # 403, 5xx, timeouts ...
                res.note(typ, u, "unreachable", r.note or str(r.status))
        if fetched_ok == 0:
            res.note(typ, base, "no_pages_found", f"no usable {typ} pages discovered or probed")
    return ats_found


def collect_ats_jobs(f, res, ats_found):
    for ats, token in ats_found.items():
        if ats == "greenhouse":
            url = f"https://boards-api.greenhouse.io/v1/boards/{token}/jobs?content=true"
            r = f.get(url, robots=False, max_bytes=8_000_000)
            if not r.ok:
                res.note("careers_ats", url, "unreachable", str(r.status)); continue
            data = json_or_none(r.text)
            if not isinstance(data, dict):
                res.note("job_posting", url, "bad_response", "greenhouse did not return JSON"); continue
            jobs = data.get("jobs", [])
            for j in jobs:
                body = BeautifulSoup(htmllib.unescape(j.get("content", "")), "html.parser").get_text("\n")
                units = [j.get("title", "")] + [s for l in body.split("\n") for s in re.split(r"(?<=[.!?])\s+", l.strip()) if s]
                for hit in find_signals(units, cap=3, min_len=0):
                    res.add_evidence(source_type="job_posting", origin="first_party", url=j.get("absolute_url", url),
                                     raw_for_hash=j.get("content", ""), title=j.get("title", ""),
                                     published=(j.get("updated_at") or "")[:10], hit=hit)
            res.note("job_posting", url, "ok", f"greenhouse board '{token}': {len(jobs)} postings scanned")
        elif ats == "lever":
            url = f"https://api.lever.co/v0/postings/{token}?mode=json"
            r = f.get(url, robots=False, max_bytes=8_000_000)
            if not r.ok:
                res.note("careers_ats", url, "unreachable", str(r.status)); continue
            jobs = json_or_none(r.text)
            if not isinstance(jobs, list):
                res.note("job_posting", url, "bad_response", "lever did not return a job list"); continue
            for j in jobs:
                units = [j.get("text", "")] + [s for l in (j.get("descriptionPlain", "") or "").split("\n")
                                               for s in re.split(r"(?<=[.!?])\s+", l.strip()) if s]
                for hit in find_signals(units, cap=3, min_len=0):
                    res.add_evidence(source_type="job_posting", origin="first_party", url=j.get("hostedUrl", url),
                                     raw_for_hash=j.get("descriptionPlain", ""), title=j.get("text", ""), hit=hit)
            res.note("job_posting", url, "ok", f"lever board '{token}': {len(jobs)} postings scanned")
        else:
            res.note("job_posting", f"{ats}:{token}", "ats_not_supported",
                     "job board detected but needs a dedicated parser / headless browser")


def collect_github(f, res, name, base):
    dom = reg_domain(urlparse(base).netloc)
    hdr = {"Accept": "application/vnd.github+json"}
    if os.getenv("GITHUB_TOKEN"):
        hdr["Authorization"] = f"Bearer {os.environ['GITHUB_TOKEN']}"
    q = quote_plus(re.sub(r"[,.]| Inc| Ltd| LLC", "", name) + " type:org")
    r = f.get(f"https://api.github.com/search/users?q={q}&per_page=5", robots=False, headers=hdr)
    if not r.ok:
        res.note("github", name, "unreachable", str(r.status)); return
    org = None
    for item in json.loads(r.text).get("items", []):
        o = f.get(f"https://api.github.com/orgs/{item['login']}", robots=False, headers=hdr)
        if o.ok and dom in (json.loads(o.text).get("blog") or "").lower():
            org = item["login"]; break  # website match avoids look-alike orgs
    if not org:
        res.note("github", name, "no_verified_org", "no org whose website field matches the vendor domain"); return
    rr = f.get(f"https://api.github.com/orgs/{org}/repos?per_page=100&type=public", robots=False, headers=hdr)
    n = 0
    repos = json.loads(rr.text) if rr.ok else []
    for repo in repos:
        text = " . ".join(filter(None, [repo.get("name"), repo.get("description"), " ".join(repo.get("topics", []))]))
        for hit in find_signals([text], cap=1):
            hit["excerpt"] = (repo.get("description") or repo["name"])[:MAX_SENTENCE_CHARS]
            res.add_evidence(source_type="developer_reference", origin="first_party", url=repo["html_url"],
                             raw_for_hash=text, title=repo["full_name"], published=(repo.get("pushed_at") or "")[:10], hit=hit)
            n += 1
    res.note("github", f"github.com/{org}", "hits" if n else "no_hits", f"{len(repos)} public repos scanned")


def collect_edgar(f, res, vid, name, contact):
    if not contact:
        res.note("sec_filing", name, "skipped", "SEC requires a contact in the User-Agent; pass --contact"); return
    hdr = {"User-Agent": f"VendorAIOSINT academic case study {contact}"}
    r = f.get("https://www.sec.gov/files/company_tickers.json", robots=False, headers=hdr, max_bytes=8_000_000)
    if not r.ok:
        res.note("sec_filing", name, "unreachable", str(r.status)); return
    target = clean_name(sec_name_override(name) or name)
    best, score = None, 0.0
    for row in json.loads(r.text).values():
        cand = clean_name(row["title"])
        sc = 1.0 if cand == target else difflib.SequenceMatcher(None, cand, target).ratio()
        if sc > score:
            best, score = row, sc
    if not best or score < 0.88:
        res.note("sec_filing", name, "not_a_registrant", f"no confident SEC match (best={score:.2f}); likely private or needs an override"); return
    cik = str(best["cik_str"]).zfill(10)
    sub = f.get(f"https://data.sec.gov/submissions/CIK{cik}.json", robots=False, headers=hdr)
    if not sub.ok:
        res.note("sec_filing", name, "unreachable", str(sub.status)); return
    rec = json.loads(sub.text)["filings"]["recent"]
    for form, acc, doc, date in zip(rec["form"], rec["accessionNumber"], rec["primaryDocument"], rec["filingDate"]):
        if form == "10-K":
            url = f"https://www.sec.gov/Archives/edgar/data/{int(cik)}/{acc.replace('-', '')}/{doc}"
            fr = f.get(url, robots=False, headers=hdr, max_bytes=15_000_000)
            if not fr.ok:
                res.note("sec_filing", url, "unreachable", str(fr.status)); return
            title, _, units = html_to_units(fr.text)
            n = 0
            for hit in find_signals(units, cap=15):
                res.add_evidence(source_type="sec_filing", origin="first_party", url=url, raw_for_hash=fr.text,
                                 title=f"{best['title']} 10-K", published=date, hit=hit); n += 1
            res.note("sec_filing", url, "hits" if n else "no_hits", f"matched '{best['title']}' (score {score:.2f}), 10-K filed {date}")
            return
    res.note("sec_filing", name, "no_10k", f"registrant '{best['title']}' has no 10-K in recent filings")


def collect_news(f, res, name):
    q = quote_plus(f'"{name}" (AI OR "artificial intelligence" OR "machine learning" OR "generative")')
    url = f"https://news.google.com/rss/search?q={q}&hl=en-US&gl=US&ceid=US:en"
    r = f.get(url, robots=False)
    if not r.ok:
        res.note("news", name, "unreachable", str(r.status)); return
    soup = BeautifulSoup(r.text, "html.parser")
    key = clean_name(name).split()[0] if clean_name(name) else name.lower()
    n = 0
    for it in soup.find_all("item")[:25]:
        title = it.find("title").get_text(strip=True) if it.find("title") else ""
        link_tag, link = it.find("link"), ""
        if link_tag is not None:
            nxt = link_tag.next_sibling
            link = (str(nxt).strip() if isinstance(nxt, NavigableString) else "") or link_tag.get_text(strip=True)
        desc = BeautifulSoup(htmllib.unescape(it.find("description").get_text()), "html.parser").get_text(" ", strip=True) if it.find("description") else ""
        if key not in (title + desc).lower():
            continue
        for hit in find_signals([title], cap=1):
            res.add_evidence(source_type="news_third_party", origin="third_party", url=link or url,
                             raw_for_hash=title + desc, title=title,
                             published=(it.find("pubdate").get_text()[:16] if it.find("pubdate") else ""), hit=hit)
            n += 1
    res.note("news", url, "hits" if n else "no_hits", "third-party; lower reliability than first-party sources")


# --------------------------------------------------------------------------- #
# 7. PUBLIC ENTRY POINT
# --------------------------------------------------------------------------- #
def scrape_vendor(vendor_id: str, name: str, website: str, level: int = 2, contact: str | None = None,
                  levels: dict | None = None, fetcher: Fetcher | None = None, overrides: dict | None = None) -> dict:
    if not name or not website:
        raise ValueError(f"vendor {vendor_id}: both a name and a website are required")
    levels = levels or DEFAULT_LEVELS
    level = int(level)
    if level not in levels:
        raise ValueError(f"level {level} not defined; available: {sorted(levels)}")
    cfg = levels[level]
    f = fetcher or Fetcher(contact=contact)
    base = norm_url(website)
    res = Result(vendor_id, name, base, level, cfg["label"])
    ov = (overrides or {}).get(vendor_id, {})

    def safe(label, fn, *args):
        try:
            return fn(*args)
        except Exception as e:  # log it as a finding and keep every other source's evidence
            res.note(label, name, "error", f"{type(e).__name__}: {e}")

    ats = safe("first_party", collect_first_party, f, res, base, cfg) or {}
    if "careers" in cfg["pages"]:
        safe("job_posting", collect_ats_jobs, f, res, ats)
    if cfg.get("github"):
        safe("github", collect_github, f, res, name, base)
    if cfg.get("edgar"):
        safe("sec_filing", collect_edgar, f, res, vendor_id, ov.get("sec_name", name), contact)
    if cfg.get("news"):
        safe("news", collect_news, f, res, name)

    ev = res.evidence
    summary = {
        "evidence_count": len(ev),
        "by_source_type": {t: sum(1 for e in ev if e["source_type"] == t) for t in sorted({e["source_type"] for e in ev})},
        "named_providers_found": sorted({p for e in ev for p in e["named_providers"]}),
        "first_party_evidence": sum(1 for e in ev if e["origin"] == "first_party"),
    }
    res.meta["finished_at"] = now_iso()
    return {"meta": res.meta, "summary": summary, "evidence": ev,
            "sources_checked": res.checked, "fetch_log": f.log[-400:]}


# --------------------------------------------------------------------------- #
# 8. WORKBOOK I/O + CLI
# --------------------------------------------------------------------------- #
FIELD_ALIASES = {  # canonical field -> accepted header spellings (case/space-insensitive)
    "id": ["vendor id", "vendor_id", "id"],
    "name": ["vendor name", "vendor_name", "name"],
    "site": ["public website", "website", "vendor website", "url", "domain"],
    "tier": ["criticality tier", "criticality", "tier"],
}
_ALIAS_TO_FIELD = {alias: fld for fld, al in FIELD_ALIASES.items() for alias in al}


def _norm(v) -> str:
    return re.sub(r"\s+", " ", str(v or "")).strip().lower()


def _field_of(v):
    return _ALIAS_TO_FIELD.get(_norm(v))


def _cell_str(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)  # Excel stores 3369 as 3369.0 sometimes
    v = str(v).strip()
    return v or None


def _parse_sheet(ws):
    """Return (vendors, problem). Supports a horizontal table (fields across a header row,
    one vendor per row) and a vertical one (fields down a column, one vendor per column)."""
    grid = [list(r) for r in ws.iter_rows(values_only=True)]
    anchor = next(((r, c) for r, row in enumerate(grid) for c, v in enumerate(row)
                   if _field_of(v) == "id"), None)
    if anchor is None:
        return [], "no 'Vendor ID' cell"
    r0, c0 = anchor
    at = lambda r, c: grid[r][c] if 0 <= r < len(grid) and 0 <= c < len(grid[r]) else None
    vendors = []

    if _field_of(at(r0, c0 + 1)) or any(_field_of(at(r0, c)) == "name" for c in range(len(grid[r0]))):
        # ---- horizontal: header row = r0
        cols = {}
        for c, v in enumerate(grid[r0]):
            fld = _field_of(v)
            if fld and fld not in cols:
                cols[fld] = c
        missing = [f for f in ("name", "site") if f not in cols]
        if missing:
            return [], f"header row has no column for {missing}"
        for row in grid[r0 + 1:]:
            get = lambda f: _cell_str(row[cols[f]]) if f in cols and cols[f] < len(row) else None
            vid, nm, site = get("id"), get("name"), get("site")
            if not vid or vid == "V-000":  # skip blanks and the fictional worked example
                continue
            vendors.append({"id": vid, "name": nm, "site": site, "tier": get("tier")})
        return vendors, None

    # ---- vertical: labels in column c0, one vendor per value column to the right
    rows = {}
    for r in range(r0, len(grid)):
        fld = _field_of(at(r, c0))
        if fld and fld not in rows:
            rows[fld] = r
    missing = [f for f in ("name", "site") if f not in rows]
    if missing:
        return [], f"label column has no row for {missing}"
    header = [_norm(v) for v in grid[r0 - 1]] if r0 > 0 else []
    for c in range(c0 + 1, max(len(r) for r in grid)):
        h = header[c] if c < len(header) else ""
        vid = _cell_str(at(rows["id"], c))
        if not vid or vid == "V-000":
            continue
        if any(w in h for w in ("description", "definition", "notes", "guidance")):
            continue  # explanatory column of a data dictionary, not a vendor
        if len(vid) > 30 and " " in vid:
            continue  # a sentence, not an ID
        vendors.append({"id": vid, "name": _cell_str(at(rows["name"], c)),
                        "site": _cell_str(at(rows["site"], c)),
                        "tier": _cell_str(at(rows["tier"], c)) if "tier" in rows else None})
    return vendors, None


def read_vendors(path: str):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    names = sorted(wb.sheetnames, key=lambda n: 0 if _norm(n) == "vendor inventory" else 1)
    problems = []
    for n in names:
        vendors, problem = _parse_sheet(wb[n])
        if problem:
            problems.append(f"'{n}': {problem}")
            continue
        usable = []
        for v in vendors:
            if not v["name"] or not v["site"]:
                print(f"warning: skipping {v['id']} - name or website is empty", file=sys.stderr)
            else:
                usable.append(v)
        if usable:
            print(f"read {len(usable)} vendor(s) from sheet '{n}'", file=sys.stderr)
            return usable
        problems.append(f"'{n}': header found but no usable vendor rows")
    raise SystemExit("Could not read vendors from " + path + ". Needs 'Vendor ID', 'Vendor Name' and "
                     "'Public Website' (optionally 'Criticality Tier'). Tried: " + "; ".join(problems))


def save_outputs(results: list[dict], out_dir: str):
    os.makedirs(out_dir, exist_ok=True)
    with open(os.path.join(out_dir, "osint_results.json"), "w", encoding="utf-8") as fh:
        json.dump(results, fh, indent=2, ensure_ascii=False)
    ev = [e for r in results for e in r["evidence"]]
    chk = [{"vendor_id": r["meta"]["vendor_id"], **c} for r in results for c in r["sources_checked"]]
    for name, rows in (("evidence_log.csv", ev), ("sources_checked.csv", chk)):
        if rows:
            with open(os.path.join(out_dir, name), "w", newline="", encoding="utf-8") as fh:
                w = csv.DictWriter(fh, fieldnames=list(rows[0].keys()))
                w.writeheader()
                for row in rows:
                    w.writerow({k: ("; ".join(v) if isinstance(v, list) else v) for k, v in row.items()})
    try:
        import pandas as pd
        with pd.ExcelWriter(os.path.join(out_dir, "Evidence_Log.xlsx")) as xw:
            pd.DataFrame([{**e, "matched_terms": "; ".join(e["matched_terms"]),
                           "named_providers": "; ".join(e["named_providers"])} for e in ev]).to_excel(xw, sheet_name="Evidence Log", index=False)
            pd.DataFrame(chk).to_excel(xw, sheet_name="Sources Checked", index=False)
            pd.DataFrame([{"vendor_id": r["meta"]["vendor_id"], "vendor": r["meta"]["vendor_name"],
                           "level": r["meta"]["level"], **{k: (", ".join(v) if isinstance(v, list) else json.dumps(v))
                           for k, v in r["summary"].items()}} for r in results]).to_excel(xw, sheet_name="Summary", index=False)
    except Exception as e:  # xlsx export is a convenience only
        print("xlsx export skipped:", e, file=sys.stderr)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--input", help="Meridian_Vendor_Input.xlsx")
    ap.add_argument("--name"); ap.add_argument("--site"); ap.add_argument("--id", default="V-ADHOC")
    ap.add_argument("--level", type=int, help="default level (used when the tier column is empty)")
    ap.add_argument("--levels", default="", help="per-vendor overrides, e.g. V-002=4,V-003=2")
    ap.add_argument("--levels-config", help="JSON file replacing the default level definitions")
    ap.add_argument("--overrides", help='JSON {"V-005": {"sec_name": "..."}} for lookups heuristics miss')
    ap.add_argument("--only", help="comma-separated vendor IDs to run")
    ap.add_argument("--contact", help="your email; goes in the User-Agent (required for SEC EDGAR)")
    ap.add_argument("--delay", type=float, default=1.0, help="seconds between requests per host")
    ap.add_argument("--out", default="osint_out")
    a = ap.parse_args()

    levels = DEFAULT_LEVELS
    if a.levels_config:
        with open(a.levels_config, encoding="utf-8") as fh:
            levels = {int(k): v for k, v in json.load(fh).items()}
    overrides = {}
    if a.overrides:
        with open(a.overrides, encoding="utf-8") as fh:
            overrides = json.load(fh)
    per = {}
    for part in a.levels.split(","):
        if part.strip():
            k, eq, val = part.partition("=")
            if not eq or not val.strip().isdigit():
                ap.error(f"--levels expects ID=LEVEL pairs like V-002=4, got '{part}'")
            per[k.strip()] = int(val)

    if a.input:
        vendors = read_vendors(a.input)
    elif a.name and a.site:
        vendors = [{"id": a.id, "name": a.name, "site": a.site, "tier": None}]
    else:
        ap.error("give --input or --name and --site")
    if a.only:
        keep = {x.strip() for x in a.only.split(",")}
        vendors = [v for v in vendors if str(v["id"]) in keep]
    if not vendors:
        sys.exit("No vendors to process.")

    fetcher = Fetcher(contact=a.contact, delay=a.delay)
    results = []
    for v in vendors:
        v["id"] = str(v["id"])
        lvl = per.get(v["id"]) or tier_to_level(v["tier"]) or a.level
        if not lvl:
            lvl = 2
            print(f"    ! {v['id']}: no tier in the workbook and no --level given; using level 2 (Standard)",
                  file=sys.stderr)
        print(f"[{v['id']}] {v['name']}  level={lvl}", flush=True)
        r = scrape_vendor(v["id"], v["name"], v["site"], int(lvl), a.contact, levels, fetcher, overrides)
        print(f"    {r['summary']['evidence_count']} evidence items | providers: {r['summary']['named_providers_found'] or 'none'}")
        results.append(r)
        save_outputs(results, a.out)  # incremental save: a crash never loses finished vendors
    print(f"Done -> {a.out}/")


if __name__ == "__main__":
    main()
