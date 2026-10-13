"""Local test fixtures: fake vendor websites served on 127.0.0.1, and a test workbook.

Used only by tests/ and tools/verify_parity.py when live websites are not reachable.
Every site and company here is fictional.
"""
from __future__ import annotations

import threading
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

import openpyxl

PAGE = "<html><head><title>{title}</title></head><body><nav>Home Products</nav><main>{body}</main></body></html>"

SITES = {
    # Optiv's worked example vendor (V-000)
    "larkspur": {
        "/": PAGE.format(title="Larkspur Servicing Group", body="""
            <h1>Loan servicing, simplified.</h1>
            <p>Larkspur Servicing Group supports lenders across the mortgage lifecycle.</p>
            <a href="/products/ai-servicing">AI servicing</a> <a href="/trust">Trust</a>"""),
        "/products/ai-servicing": PAGE.format(title="AI Servicing Assistant", body="""
            <p>Our servicing platform uses a large language model to draft borrower correspondence and summarise call notes.</p>
            <p>The assistant is built on Azure OpenAI and is generally available to all servicing clients today.</p>
            <p>Agents review every draft before it is sent.</p>"""),
        "/trust": PAGE.format(title="Trust Center", body="""
            <p>We protect borrower data with encryption at rest and in transit.</p>
            <p>Sub-processors include Microsoft (Azure OpenAI Service) for generative AI features.</p>"""),
    },
    "northwind": {
        "/": PAGE.format(title="Northwind Payments", body="""
            <h1>Payments infrastructure for banks.</h1>
            <p>Northwind Payments processes card and ACH transactions for regional banks.</p>
            <a href="/platform/fraud-intelligence">Fraud intelligence</a>
            <a href="/blog/future-of-payments">Blog</a> <a href="/careers">Careers</a>"""),
        "/platform/fraud-intelligence": PAGE.format(title="Fraud Intelligence", body="""
            <p>Our fraud platform uses machine learning models trained on billions of transactions to score every payment in real time.</p>
            <p>The anomaly detection engine is in production at over 200 banks and is now available in the EU.</p>
            <p>Analysts see the top contributing features for each machine learning decision.</p>"""),
        "/blog/future-of-payments": PAGE.format(title="The future of payments", body="""
            <p>We believe generative AI will revolutionize how banks talk to customers in the next decade.</p>
            <p>Join our webinar on the future of AI in payments.</p>"""),
        "/careers": PAGE.format(title="Careers", body="""
            <p>We are hiring a Senior Machine Learning Engineer to build fraud models.</p>"""),
    },
    "bluebell": {
        "/": PAGE.format(title="Bluebell Docs", body="""
            <h1>Document workflows for regulated teams.</h1>
            <p>Bluebell Docs routes, signs and archives documents using business rules you configure.</p>
            <a href="/news/ai-roadmap">News</a>"""),
        "/news/ai-roadmap": PAGE.format(title="Our AI roadmap", body="""
            <p>Bluebell plans to explore generative AI features on our roadmap for next year.</p>
            <p>Our workflow automation is rule-based and deterministic, with no machine learning involved today.</p>"""),
    },
    "cobalt": {
        "/": PAGE.format(title="Cobalt Advisory", body="""
            <h1>Regulatory advisory for financial institutions.</h1>
            <p>Cobalt Advisory provides consulting on compliance programmes and audit readiness.</p>
            <p>Our partners have decades of experience with bank examinations.</p>"""),
    },
    "robots": {
        "/robots.txt": "User-agent: *\nDisallow: /\n",
        "/": PAGE.format(title="Private", body="<p>Machine learning powered reporting for every team.</p>"),
    },
}


def _handler(site: dict):
    class H(BaseHTTPRequestHandler):
        def do_GET(self):  # noqa: N802
            path = self.path.split("?")[0]
            body = site.get(path)
            if body is None:
                self.send_response(404)
                self.send_header("Content-Type", "text/html; charset=utf-8")
                self.end_headers()
                self.wfile.write(b"<html><body>Not found</body></html>")
                return
            self.send_response(200)
            ctype = "text/plain" if path.endswith(".txt") else "text/html; charset=utf-8"
            self.send_header("Content-Type", ctype)
            self.end_headers()
            self.wfile.write(body.encode("utf-8"))

        def log_message(self, *a):
            pass
    return H


def start_sites(base_port: int = 18601) -> dict[str, str]:
    """Start one local server per fictional site. Returns name -> base URL."""
    urls = {}
    for i, (name, site) in enumerate(SITES.items()):
        srv = ThreadingHTTPServer(("127.0.0.1", base_port + i), _handler(site))
        threading.Thread(target=srv.serve_forever, daemon=True).start()
        urls[name] = f"http://127.0.0.1:{base_port + i}"
    urls["unreachable"] = "http://127.0.0.1:1"   # nothing listens here
    return urls


OUTPUT_HEADERS = [
    "Assessment Depth Applied", "AI Usage Detected (Y/N)",
    "Evidence Description along with OSINT Source(s) / Evidence Link", "How Vendor May Be Using AI",
    "Named AI Sub-processors / Model Providers", "AI Security Risk Classification", "Risk Rationale",
    "Recommended Action", "Assessedd By / Date",
]
INPUT_HEADERS = ["Vendor ID", "Vendor Name", "Public Website", "Category", "Operational Dependency",
                 "Data Classification Accessed", "Data Volume (annual)"]


def build_workbook(path: str, urls: dict[str, str], n_vendors: int | None = None) -> str:
    """Workbook in the notebook's expected layout (sheet 'Test Vendor Input')."""
    vendors = [
        ["V-000", "Larkspur Servicing Group", urls["larkspur"], "Core Platform", "Total",
         "Highly Confidential / PII", 250_000_000],
        ["V-001", "Northwind Payments", urls["northwind"], "Payments", "High", "Confidential", "12M"],
        ["V-002", "Bluebell Docs", urls["bluebell"], "Document and Communications", "Low", "Internal", 500_000],
        ["V-003", "Cobalt Advisory", urls["cobalt"], "Professional Services", "Moderate", "Internal", "1,500,000"],
        ["V-004", "Gone Dark Ltd", urls["unreachable"], "Data and Reporting", "Moderate", "Confidential", 2_000_000],
        ["V-005", "Robots Only Inc", urls["robots"], "Data and Reporting", "Low", "N/A", None],
        ["V-006", "No Website Co", None, "Payments", "High", "Confidential", 5_000_000],
    ]
    if n_vendors is not None:
        vendors = vendors[:n_vendors]
    v000_prefill = ["Depth configured to Critical Tier", "Yes (Example)",
                    "Example evidence pre-filled by Optiv", "Example: LLM drafting of borrower letters",
                    "Azure OpenAI", "Critical", "Example rationale", "Example action", "Optiv / 2026-01-01"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Instructions"
    ws["A1"] = "Read me: complete the 'Test Vendor Input' sheet."
    ws2 = wb.create_sheet("Test Vendor Input")
    ws2["A1"] = "Meridian Financial - Third-Party AI Vendor Assessment (test fixture, fictional data)"
    ws2["A2"] = "Fill in columns H onward. Row V-000 is a worked example."
    ws2.append([])
    ws2.append(INPUT_HEADERS + OUTPUT_HEADERS)
    for v in vendors:
        ws2.append(v + (v000_prefill if v[0] == "V-000" else [None] * len(OUTPUT_HEADERS)))
    wb.save(path)
    return path
