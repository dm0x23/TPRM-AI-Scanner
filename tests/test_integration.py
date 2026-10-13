"""End-to-end checks of the adapter, as the Streamlit app calls it.

Runs against local fictional websites (tests/fixtures.py). Uses TEST-DOUBLE models
only where noted, because the real models need a Hugging Face download.

    python tests/test_integration.py            # plain Python, no pytest needed
"""
from __future__ import annotations

import io
import sys
import tempfile
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import openpyxl  # noqa: E402
import pandas as pd  # noqa: E402

from models.schema import RESULT_COLUMNS, validate_results  # noqa: E402
from services import excel_io, pipeline_adapter  # noqa: E402
from services.optiv_pipeline import orchestrator  # noqa: E402
from tests.fixtures import build_workbook, start_sites  # noqa: E402
from tests.stub_models import stub_model_factory  # noqa: E402

URLS = start_sites(base_port=18701)
PASSED = []


def load_like_app(path):
    data = Path(path).read_bytes()
    sheets = excel_io.read_workbook_raw(data, Path(path).name)
    raw = sheets["Test Vendor Input"]
    hdr = excel_io.detect_header_row(raw)
    return excel_io.build_table(raw, hdr), excel_io.build_table(raw, hdr, drop_blank_columns=False)


def check(name, cond, detail=""):
    if not cond:
        raise AssertionError(f"{name} FAILED {detail}")
    PASSED.append(name)
    print(f"  ok  {name}")


def report_bytes(out):
    run = out.run
    extra = {"Completed Assessment": run.export_df, "Evidence (all rows)": run.nlp_df,
             "Sources Checked": pipeline_adapter.sources_checked_table(run)}
    if run.v000_comparison is not None:
        extra["V-000 Regression Check"] = run.v000_comparison
    return excel_io.build_report(validate_results(out.results), False, extra)


def main():
    tmp = Path(tempfile.mkdtemp())
    real_missing = pipeline_adapter.missing_dependencies
    real_factory = orchestrator.default_model_factory

    # ---- A. dependency gate --------------------------------------------------
    print("A. dependency gate")
    wb3 = build_workbook(str(tmp / "three.xlsx"), URLS, n_vendors=3)
    preview, sheet_table = load_like_app(wb3)
    from tests.fixtures import INPUT_HEADERS, OUTPUT_HEADERS
    check("sheet table keeps every header column", list(sheet_table.columns) == INPUT_HEADERS + OUTPUT_HEADERS)
    raw = excel_io.read_workbook_raw(Path(wb3).read_bytes(), "x.xlsx")["Test Vendor Input"]
    no_v000 = raw.drop(index=4)   # without the pre-filled example the output columns are empty
    check("preview drops empty output columns, sheet table keeps them",
          "Risk Rationale" not in excel_io.build_table(no_v000, 3).columns
          and "Risk Rationale" in excel_io.build_table(no_v000, 3, drop_blank_columns=False).columns)
    if real_missing():
        try:
            pipeline_adapter.run_assessment_full(preview, sheet_table=sheet_table)
            check("missing packages raise PipelineNotConnectedError", False)
        except pipeline_adapter.PipelineNotConnectedError as e:
            check("missing packages raise PipelineNotConnectedError", "torch" in str(e) or "transformers" in str(e), str(e))
    pipeline_adapter.missing_dependencies = lambda: []   # let the run start in this sandbox

    # ---- B. real model loading fails -> honest incomplete ---------------------
    print("B. NLP model failure (real loader, no torch/HF access here)")
    t = time.time()
    out = pipeline_adapter.run_assessment_full(preview, sheet_table=sheet_table)
    res = validate_results(out.results)
    print(f"     nlp_error = {out.run.nlp_error}  ({time.time() - t:.0f}s)")
    check("model failure is reported", out.run.nlp_error is not None)
    with_ev = [v for v, s in out.run.vendor_status.items() if s.evidence_rows > 0]
    check("vendors with evidence get no tier when NLP failed",
          res.set_index("Vendor ID").loc[with_ev, "Risk Tier"].isna().all(), str(res[["Vendor ID", "Risk Tier"]]))
    check("their status says why", all("NLP verification failed" in out.run.vendor_status[v].reason for v in with_ev))
    check("V-000 template row untouched",
          out.run.export_df.loc[out.run.export_df["Vendor ID"] == "V-000", "AI Security Risk Classification"].iloc[0] == "Critical")

    # ---- C. 3 vendors, complete run (TEST-DOUBLE models) ---------------------
    print("C. 3 vendors with TEST-DOUBLE models")
    orchestrator.default_model_factory = stub_model_factory
    steps = []
    out = pipeline_adapter.run_assessment_full(preview, progress_callback=lambda f, m: steps.append((f, m)),
                                               sheet_table=sheet_table)
    res = validate_results(out.results)
    check("results have dashboard columns", all(c in res.columns for c in RESULT_COLUMNS))
    check("3 rows, vendor IDs preserved", list(res["Vendor ID"]) == ["V-000", "V-001", "V-002"])
    check("all complete", (res["Assessment Status"] == "Complete").all(), str(res["Assessment Notes"].tolist()))
    check("progress monotone 0->1", steps[0][0] == 0.0 and steps[-1][0] == 1.0
          and all(a[0] <= b[0] + 1e-9 for a, b in zip(steps, steps[1:])))
    check("criticality V-000=4 V-001=3 V-002=1", list(res["Criticality"]) == [4, 3, 1])
    check("evidence has source URL", all(str(u).startswith("http://127.0.0.1") for u in res["Source URLs"]))
    check("V-000 comparison produced", out.run.v000_comparison is not None
          and list(out.run.v000_comparison["Optiv Pre-Filled Answer"]) == ["Yes (Example)", "Critical", "Azure OpenAI"])
    xb = report_bytes(out)
    wbk = openpyxl.load_workbook(io.BytesIO(xb))
    check("report sheets", wbk.sheetnames == ["Summary", "Assessment Results", "Completed Assessment",
                                              "Evidence (all rows)", "Sources Checked", "V-000 Regression Check"],
          str(wbk.sheetnames))
    ca = pd.read_excel(io.BytesIO(xb), sheet_name="Completed Assessment", dtype=str)
    exp = out.run.export_df.astype(str).replace({"nan": None, "None": None})
    for col in ["Vendor ID", "AI Usage Detected (Y/N)", "AI Security Risk Classification", "Recommended Action",
                "Evidence Description along with OSINT Source(s) / Evidence Link"]:
        check(f"Excel 'Completed Assessment' {col} == pipeline output",
              ca[col].fillna("").tolist() == exp[col].fillna("").tolist())
    ar = pd.read_excel(io.BytesIO(xb), sheet_name="Assessment Results", dtype=str)
    check("Excel 'Assessment Results' tiers == dashboard", ar["Risk Tier"].fillna("").tolist()
          == res["Risk Tier"].fillna("").tolist())
    check("no internal _columns in template", not any(str(c).startswith("_") for c in ca.columns))

    # ---- D. full fixture workbook (7 vendors), real 1 s fetch delay -----------
    print("D. full workflow, 7 vendors, real Fetcher settings")
    wb7 = build_workbook(str(tmp / "seven.xlsx"), URLS)
    preview, sheet_table = load_like_app(wb7)
    t = time.time()
    out = pipeline_adapter.run_assessment_full(preview, sheet_table=sheet_table)
    res = validate_results(out.results)
    print(f"     {time.time() - t:.0f}s")
    print(res[["Vendor ID", "Criticality", "AI Usage Detected (Y/N)", "Risk Tier", "Assessment Status"]].to_string())
    tiers = dict(zip(res["Vendor ID"], res["Risk Tier"]))
    check("unreachable / robots / no-website vendors have no tier",
          all(pd.isna(tiers[v]) for v in ("V-004", "V-005", "V-006")))
    check("reachable no-AI vendor keeps notebook's Low", tiers["V-003"] == "Low")
    check("tiers are only allowed values", set(t for t in tiers.values() if not pd.isna(t)) <= {"Low", "Medium", "High", "Critical"})
    xb = report_bytes(out)
    summ = pd.read_excel(io.BytesIO(xb), sheet_name="Summary", header=None, dtype=str)
    check("summary counts incomplete vendors",
          "3" in summ[summ[0] == "Vendors without a risk tier (incomplete)"][1].tolist())

    # ---- E. bad input ---------------------------------------------------------
    print("E. input errors")
    try:
        pipeline_adapter.run_assessment_full(preview.drop(columns=["Operational Dependency"]))
        check("missing required column rejected", False)
    except pipeline_adapter.PipelineInputError as e:
        check("missing required column rejected", "Operational Dependency" in str(e), str(e))
    renamed = sheet_table.rename(columns={"Public Website": "website ", "Data Volume (annual)": "DATA VOLUME",
                                          "Assessedd By / Date": "Assessed By / Date"})
    m_in = orchestrator.resolve_input_columns(renamed)
    m_out = orchestrator.resolve_export_columns(renamed)
    check("header variants resolved", m_in["Public Website"] == "website " and m_in["Data Volume (annual)"] == "DATA VOLUME")
    check("'Assessed By / Date' header used for the date", m_out["Assessed By / Date"] == "Assessed By / Date")

    pipeline_adapter.missing_dependencies = real_missing
    orchestrator.default_model_factory = real_factory
    print(f"\nALL {len(PASSED)} CHECKS PASSED")


if __name__ == "__main__":
    main()
