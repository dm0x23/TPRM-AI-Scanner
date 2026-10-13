"""Reading uploaded workbooks and writing the downloadable report."""
from __future__ import annotations

import io
from datetime import date, datetime, timezone

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from models.schema import DEMO_LABEL, RISK_TIERS, TIER_COLORS, tier_counts

MAX_FILE_BYTES = 10 * 1024 * 1024  # 10 MB
MAX_SCAN_ROWS = 40  # how far down we look for the header row

# Words that commonly appear in a vendor sheet's column headers.
HEADER_HINTS = (
    "vendor", "supplier", "name", "domain", "website", "url", "criticality",
    "category", "spend", "service", "contact", "owner", "data", "description",
    "contract", "country", "industry",
)

NAVY = "0B1426"
MUTED_BLUE = "5B8DEF"


class ExcelReadError(Exception):
    """Raised with a user-friendly message when a workbook cannot be used."""


# --------------------------------------------------------------------------
# Reading
# --------------------------------------------------------------------------
def _cell_text(value) -> str:
    """Turn any cell value into clean text ('' for empty cells)."""
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_workbook_raw(data: bytes, filename: str) -> dict[str, pd.DataFrame]:
    """Read every sheet WITHOUT treating any row as the header."""
    if not filename.lower().endswith(".xlsx"):
        raise ExcelReadError("Please upload an Excel file with the .xlsx extension.")
    if not data:
        raise ExcelReadError("The uploaded file is empty.")
    if len(data) > MAX_FILE_BYTES:
        raise ExcelReadError("The file is larger than 10 MB. Please upload a smaller workbook.")
    try:
        sheets = pd.read_excel(
            io.BytesIO(data), sheet_name=None, header=None, dtype=object, engine="openpyxl"
        )
    except Exception as exc:  # corrupt zip, password-protected, wrong format...
        raise ExcelReadError(
            "Could not open this file. It may be corrupted, password-protected, "
            "or not a genuine .xlsx workbook."
        ) from exc
    sheets = {n: df for n, df in sheets.items() if not df.dropna(how="all").empty}
    if not sheets:
        raise ExcelReadError("The workbook has no sheets containing data.")
    return sheets


def detect_header_row(raw: pd.DataFrame) -> int | None:
    """Guess which row (0-based) holds the column headers.

    Instruction text is usually one long cell, so rows with fewer than two
    filled cells are skipped. Rows containing vendor-style words score higher.
    Returns None when nothing looks like a header (user can pick manually).
    """
    best_row, best_score = None, 0
    limit = min(len(raw), MAX_SCAN_ROWS)
    for i in range(limit):
        cells = [_cell_text(v) for v in raw.iloc[i]]
        filled = [c for c in cells if c]
        if len(filled) < 2:
            continue
        hits = sum(any(h in c.lower() for h in HEADER_HINTS) for c in filled)
        short = sum(len(c) <= 40 for c in filled)  # headers are short labels
        score = hits * 3 + short
        if i + 1 < len(raw) and any(_cell_text(v) for v in raw.iloc[i + 1]):
            score += 2  # data follows this row
        if score > best_score:
            best_row, best_score = i, score
    return best_row


def instruction_lines(raw: pd.DataFrame, header_row: int) -> list[str]:
    """Text found above the header row (instructions / notes)."""
    lines = []
    for i in range(min(header_row, len(raw))):
        text = " | ".join(c for c in (_cell_text(v) for v in raw.iloc[i]) if c)
        if text:
            lines.append(text)
    return lines


def build_table(raw: pd.DataFrame, header_row: int, drop_blank_columns: bool = True) -> pd.DataFrame:
    """Use `header_row` (0-based) as headers and everything below as data.

    With drop_blank_columns=False, empty columns are kept (used for the completed
    assessment template, whose output columns are often empty before a run).
    """
    if header_row < 0 or header_row >= len(raw):
        raise ExcelReadError("The chosen header row is outside the sheet.")

    names, seen = [], {}
    for i, value in enumerate(raw.iloc[header_row]):
        name = _cell_text(value) or f"Column {i + 1}"
        seen[name] = seen.get(name, 0) + 1
        names.append(name if seen[name] == 1 else f"{name} ({seen[name]})")

    body = raw.iloc[header_row + 1:].copy()
    body.columns = names
    body = body.map(_cell_text)
    body = body.loc[~(body == "").all(axis=1)]      # drop blank rows
    if drop_blank_columns:
        body = body.loc[:, ~(body == "").all(axis=0)]   # drop blank columns
    else:
        # keep named empty columns, but never invent headers for unnamed empty ones
        body = body.loc[:, [not (c.startswith("Column ") and (body[c] == "").all()) for c in body.columns]]
    body = body.reset_index(drop=True)
    if body.empty:
        raise ExcelReadError(
            "No data rows were found below the chosen header row. Try a different header row."
        )
    return body


# --------------------------------------------------------------------------
# Writing
# --------------------------------------------------------------------------
def _put(ws, row: int, col: int, value):
    """Write a value as plain text, never as an Excel formula."""
    text = "" if value is None or (isinstance(value, float) and pd.isna(value)) else value
    cell = ws.cell(row=row, column=col, value=text)
    if isinstance(text, str):
        cell.data_type = "s"  # blocks formula injection from '=...' values
    return cell


def _write_plain_sheet(wb, title: str, df: pd.DataFrame, border, head_fill, head_font) -> None:
    ws = wb.create_sheet(title[:31])
    for j, col_name in enumerate(df.columns, start=1):
        c = _put(ws, 1, j, str(col_name))
        c.fill, c.font, c.border = head_fill, head_font, border
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = 22
    for i, row in enumerate(df.itertuples(index=False), start=2):
        for j, value in enumerate(row, start=1):
            if isinstance(value, (list, dict)):
                value = str(value)
            c = _put(ws, i, j, value)
            c.alignment = Alignment(vertical="top", wrap_text=False)
    ws.freeze_panes = "A2"
    if len(df) and len(df.columns):
        ws.auto_filter.ref = f"A1:{get_column_letter(len(df.columns))}{len(df) + 1}"


def build_report(results: pd.DataFrame, is_demo: bool,
                 extra_sheets: dict[str, pd.DataFrame] | None = None) -> bytes:
    """Create the .xlsx assessment report and return it as bytes (in memory).

    extra_sheets (name -> DataFrame) are appended after the results sheet; the real
    pipeline uses this for the completed Meridian template and the evidence log.
    """
    df = results.copy()
    if is_demo and "Data Label" not in df.columns:
        df["Data Label"] = DEMO_LABEL

    thin = Side(style="thin", color="C9D3E3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor=NAVY)
    head_font = Font(bold=True, color="FFFFFF")
    demo_fill = PatternFill("solid", fgColor="C0392B")
    demo_font = Font(bold=True, color="FFFFFF", size=12)

    wb = Workbook()

    # ---- Summary sheet ----
    ws = wb.active
    ws.title = "Summary"
    _put(ws, 1, 1, "Meridian AI Vendor Risk Assessment").font = Font(bold=True, size=16, color=NAVY)
    r = 2
    if is_demo:
        c = _put(ws, r, 1, DEMO_LABEL)
        c.fill, c.font = demo_fill, demo_font
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        r += 1
    r += 1
    info = [
        ("Generated (UTC)", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M")),
        ("Data source", "Synthetic demo data" if is_demo else "Assessment pipeline output"),
        ("Vendors in report", len(df)),
    ]
    if not is_demo and "Risk Tier" in df.columns:
        info.append(("Vendors without a risk tier (incomplete)", int(df["Risk Tier"].isna().sum())))
    for label, value in info:
        _put(ws, r, 1, label).font = Font(bold=True)
        _put(ws, r, 2, value)
        r += 1
    r += 1
    for col, text in enumerate(["Risk Tier", "Vendors"], start=1):
        c = _put(ws, r, col, text)
        c.fill, c.font, c.border = head_fill, head_font, border
    r += 1
    for tier, n in tier_counts(df).items():
        c = _put(ws, r, 1, tier)
        c.fill = PatternFill("solid", fgColor=TIER_COLORS[tier].lstrip("#"))
        c.font = Font(bold=True, color="FFFFFF")
        c.border = border
        _put(ws, r, 2, n).border = border
        r += 1
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 20

    # ---- Results sheet ----
    wr = wb.create_sheet("Assessment Results")
    header_row = 1
    if is_demo:
        c = _put(wr, 1, 1, DEMO_LABEL)
        c.fill, c.font = demo_fill, demo_font
        wr.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(df.columns), 1))
        header_row = 2
        wr.sheet_properties.tabColor = "C0392B"
        wr.oddHeader.center.text = DEMO_LABEL
        ws.sheet_properties.tabColor = "C0392B"

    wide = {"AI Usage Evidence": 50, "Source URLs": 45, "Sub-processors": 32, "Recommended Action": 48}
    for j, col_name in enumerate(df.columns, start=1):
        c = _put(wr, header_row, j, col_name)
        c.fill, c.font, c.border = head_fill, head_font, border
        c.alignment = Alignment(vertical="center", wrap_text=True)
        wr.column_dimensions[get_column_letter(j)].width = wide.get(col_name, 18)

    for i, row in enumerate(df.itertuples(index=False), start=header_row + 1):
        for j, value in enumerate(row, start=1):
            c = _put(wr, i, j, value)
            c.border = border
            c.alignment = Alignment(vertical="top", wrap_text=True)
            if df.columns[j - 1] == "Risk Tier" and value in TIER_COLORS:
                c.fill = PatternFill("solid", fgColor=TIER_COLORS[value].lstrip("#"))
                c.font = Font(bold=True, color="FFFFFF")

    wr.freeze_panes = wr.cell(row=header_row + 1, column=1)
    if len(df):
        wr.auto_filter.ref = (
            f"A{header_row}:{get_column_letter(len(df.columns))}{header_row + len(df)}"
        )

    for title, extra in (extra_sheets or {}).items():
        if extra is not None:
            _write_plain_sheet(wb, title, extra, border, head_fill, head_font)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
